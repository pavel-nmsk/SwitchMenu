# -*- coding: utf-8 -*-
import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext
import configparser
import shutil

# ИМПОРТИРУЕМ ВШИТЫЕ РЕСУРСЫ
from resources import get_icon_path, get_photo_path, get_background_path

# ПРОВЕРКА БИБЛИОТЕК
try:
    from PIL import Image, ImageTk, ImageDraw
    HAS_PIL = True
    print("✅ Pillow загружен")
except ImportError as e:
    HAS_PIL = False
    print(f"❌ Pillow НЕ загружен: {e}")

try:
    import cv2
    HAS_CV2 = True
    print("✅ OpenCV загружен (превью видео и проигрывание)")
except ImportError:
    HAS_CV2 = False
    print("⚠️ OpenCV не найден — видео не будет проигрываться")

try:
    from docx import Document
    HAS_DOCX = True
    print("✅ python-docx загружен")
except ImportError as e:
    HAS_DOCX = False
    print(f"❌ python-docx НЕ загружен: {e}")

# Поддерживаемые форматы изображений и видео
IMAGE_EXTENSIONS = ('.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff', '.tif', '.webp', '.ico')
VIDEO_EXTENSIONS = ('.mp4', '.mkv', '.avi', '.mov', '.webm', '.wmv', '.m4v')


class SwitchMenu:
    def __init__(self, root):
        self.root = root
        self.root.title("🎮 Меню игр Switch")
        self.root.geometry("1200x700")
        self.root.minsize(1000, 500)
        self.root.configure(bg='#1a1a2e')
        
        # ИКОНКА ИЗ ВШИТОГО РЕСУРСА
        try:
            icon_path = get_icon_path()
            if icon_path and os.path.exists(icon_path):
                self.root.iconbitmap(icon_path)
                print("✅ Иконка загружена из вшитого ресурса")
            else:
                print("⚠️ Иконка не загружена")
        except Exception as e:
            print(f"⚠️ Ошибка загрузки иконки: {e}")
        
        # ФОН ИЗ ВШИТОГО РЕСУРСА
        self.background_image = None
        self.background_photo = None
        self.background_tiles = []
        self.load_background()
        
        # НАСТРОЙКИ (в папке пользователя)
        self.ini_path = os.path.join(os.path.expanduser("~"), 'switch_menu.ini')
        self.config = configparser.ConfigParser()
        self.load_config()
        
        # Загружаем сохранённый путь
        self.switch_path = self.config.get('Settings', 'last_path', fallback='')
        self.games = []
        self.current_game = None
        self.current_desc_path = None
        
        # Медиа (обложки и геймплей)
        self.current_media = {'cover': [], 'gameplay': []}
        self._displayed_media_path = {'cover': None, 'gameplay': None}
        self.slides = {
            'cover':    {'index': 0, 'timer': None, 'is_sliding': False, 'is_hovering': False},
            'gameplay': {'index': 0, 'timer': None, 'is_sliding': False, 'is_hovering': False},
        }
        self.slide_interval = self.config.getint('Settings', 'slide_interval', fallback=3000)
        self.play_video_full = self.config.getboolean('Settings', 'play_video_full', fallback=False)
        
        # Для бегущей строки
        self.title_scroll_timer = None
        self.title_scroll_text = None
        self.title_scroll_pos = 0
        
        # Для отложенного обновления правой панели
        self._update_job = None
        
        # Для навигации с клавиатуры
        self.selected_index = -1
        
        # --- ВИДЕО ПЛЕЕРЫ ---
        self.video_players = {
            'cover': {
                'cap': None,
                'timer': None,
                'running': False,
                'current_path': None,
                'delay': 30,
                'label': None,
            },
            'gameplay': {
                'cap': None,
                'timer': None,
                'running': False,
                'current_path': None,
                'delay': 30,
                'label': None,
            }
        }
        
        # Canvas для фона
        self.canvas = tk.Canvas(self.root, highlightthickness=0, bg='#1a1a2e')
        self.canvas.pack(fill='both', expand=True)
        self.draw_background()
        
        # Создаём виджеты поверх Canvas
        self.create_widgets()
        
        # Сохраняем ссылки на label'ы в плеерах
        self.video_players['cover']['label'] = self.cover_label
        self.video_players['gameplay']['label'] = self.gameplay_label
        
        # Статус библиотек
        status = "✅" if HAS_DOCX and HAS_PIL else "⚠️"
        self.status_label.config(text=f"{status} docx: {HAS_DOCX}, PIL: {HAS_PIL}")
        
        # Привязываем клавиши навигации
        self.root.bind('<Up>', self.on_key_up)
        self.root.bind('<Down>', self.on_key_down)
        self.root.focus_set()
        
        # Запускаем сканирование
        self.root.after(500, self.auto_start)
        
        # Привязываем событие изменения размера окна
        self.root.bind('<Configure>', self.on_resize)
        self._layout_resize_job = None
        self.canvas.bind('<Configure>', self.on_layout_resize)
        
        # При закрытии окна чистим плееры
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
    
    def on_closing(self):
        self._stop_all_video()
        self.root.destroy()
    
    # ------------------------------------------------------------
    # НАВИГАЦИЯ С КЛАВИАТУРЫ
    # ------------------------------------------------------------
    def on_key_up(self, event):
        if self.games:
            new_idx = max(0, self.selected_index - 1)
            if new_idx != self.selected_index:
                self.select_game(new_idx)
        return "break"
    
    def on_key_down(self, event):
        if self.games:
            new_idx = min(len(self.games) - 1, self.selected_index + 1)
            if new_idx != self.selected_index:
                self.select_game(new_idx)
        return "break"
    
    def select_game(self, index):
        if 0 <= index < len(self.games):
            self.selected_index = index
            game = self.games[index]
            game_path = os.path.join(self.switch_path, game)
            desc_file, cover_media, gameplay_media = self.find_game_files(game_path)
            self.highlight_game(game)
            self.scroll_to_index(index)
            self.show_game_description(game, desc_file, cover_media, gameplay_media, delay=300)
    
    def scroll_to_index(self, index):
        children = self.scrollable_frame.winfo_children()
        if not children or index < 0 or index >= len(children):
            return
        frame = children[index]
        self.scrollable_frame.update_idletasks()
        y = frame.winfo_y()
        height = frame.winfo_height()
        canvas_height = self.list_canvas.winfo_height()
        margin = 10
        visible_top = self.list_canvas.canvasy(0)
        visible_bottom = visible_top + canvas_height
        if visible_top + margin <= y and y + height <= visible_bottom - margin:
            return
        total_height = self.scrollable_frame.winfo_height()
        if total_height <= canvas_height:
            self.list_canvas.yview_moveto(0)
            return
        if y < visible_top + margin:
            fraction = (y - margin) / (total_height - canvas_height)
        else:
            fraction = (y + height - canvas_height + margin) / (total_height - canvas_height)
        fraction = max(0.0, min(1.0, fraction))
        self.list_canvas.yview_moveto(fraction)
    
    # ------------------------------------------------------------
    # ВИДЕО ПЛЕЙБЕК (автопроигрывание без звука)
    # ------------------------------------------------------------
    def _stop_video_playback(self, role):
        player = self.video_players.get(role)
        if not player:
            return
        player['running'] = False
        if player['timer']:
            self.root.after_cancel(player['timer'])
            player['timer'] = None
        if player['cap'] is not None:
            player['cap'].release()
            player['cap'] = None
        player['current_path'] = None

    def _stop_all_video(self):
        for role in ('cover', 'gameplay'):
            self._stop_video_playback(role)

    def _start_video_playback(self, role, file_path):
        if not HAS_CV2:
            self._show_video_placeholder(role, file_path)
            return

        self._stop_video_playback(role)
        player = self.video_players[role]
        label = player['label']

        cap = cv2.VideoCapture(file_path)
        if not cap.isOpened():
            label.config(text="❌ Видео не открыть", image='', fg='#ff6666')
            label.image = None
            player['current_path'] = None
            return

        fps = cap.get(cv2.CAP_PROP_FPS)
        if fps <= 0:
            fps = 30
        delay = int(1000 / fps)

        player['cap'] = cap
        player['running'] = True
        player['current_path'] = file_path
        player['delay'] = delay

        self._update_video_frame(role)

    def _update_video_frame(self, role):
        player = self.video_players.get(role)
        if not player or not player['running']:
            return

        cap = player['cap']
        if cap is None:
            return

        ret, frame = cap.read()
        if not ret:
            if self.play_video_full:
                self._stop_video_playback(role)
                st = self.slides.get(role)
                if st and st['is_sliding'] and not st['is_hovering']:
                    self.next_slide(role)
                return
            else:
                cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                ret, frame = cap.read()
                if not ret:
                    self._stop_video_playback(role)
                    player['label'].config(text="❌ Ошибка видео", image='', fg='#ff6666')
                    player['label'].image = None
                    return

        try:
            frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            img = Image.fromarray(frame_rgb)
        except Exception as e:
            print(f"Ошибка конвертации кадра: {e}")
            self._stop_video_playback(role)
            player['label'].config(text="❌ Ошибка кадра", image='', fg='#ff6666')
            player['label'].image = None
            return

        max_w, max_h = 250, 200
        img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)

        try:
            photo = ImageTk.PhotoImage(img)
        except Exception as e:
            print(f"Ошибка создания PhotoImage: {e}")
            self._stop_video_playback(role)
            player['label'].config(text="❌ Ошибка отображения", image='', fg='#ff6666')
            player['label'].image = None
            return

        label = player['label']
        label.config(image=photo, text='', fg='#ffffff')
        label.image = photo

        if player['running']:
            player['timer'] = self.root.after(player['delay'], self._update_video_frame, role)

    def _show_video_placeholder(self, role, file_path):
        player = self.video_players[role]
        label = player['label']
        name = os.path.basename(file_path)
        label.config(text=f"🎬 {name}\n(клик — открыть)", image='', fg='#cccccc')
        label.image = None
        player['current_path'] = file_path

    def render_media(self, role, item):
        label = self.cover_label if role == 'cover' else self.gameplay_label
        empty_text = "Обложка отсутствует" if role == 'cover' else "Геймплей отсутствует"

        self._stop_video_playback(role)

        if not item:
            label.config(text=empty_text, image='', fg='#888888')
            label.image = None
            self._displayed_media_path[role] = None
            return

        path = item['path']
        self._displayed_media_path[role] = path

        if item['is_video']:
            self._start_video_playback(role, path)
        else:
            img = self.load_image(path, 250, 200)
            if img:
                label.config(image=img, text='', fg='#ffffff')
                label.image = img
            else:
                label.config(text="❌ Не загружено", image='', fg='#ff6666')
                label.image = None

    # ------------------------------------------------------------
    # ВСПОМОГАТЕЛЬНЫЕ МЕТОДЫ
    # ------------------------------------------------------------
    def _truncate_name(self, text, max_chars=34):
        if len(text) <= max_chars:
            return text
        return text[:max_chars - 1].rstrip() + '…'

    def load_config(self):
        if os.path.exists(self.ini_path):
            try:
                self.config.read(self.ini_path, encoding='utf-8')
                print(f"📁 Настройки загружены из {self.ini_path}")
                return
            except Exception as e:
                print(f"⚠️ Ошибка чтения настроек: {e}")
        
        self.config['Settings'] = {
            'auto_scan': 'True',
            'last_path': '',
            'slide_interval': '3000',
            'play_video_full': 'False'
        }
        self.save_config()
        print(f"📁 Создан новый файл настроек: {self.ini_path}")
    
    def save_config(self):
        try:
            with open(self.ini_path, 'w', encoding='utf-8') as f:
                self.config.write(f)
            print(f"💾 Настройки сохранены в {self.ini_path}")
        except Exception as e:
            print(f"❌ Ошибка сохранения настроек: {e}")
    
    def load_background(self):
        try:
            fon_path = get_background_path()
            if fon_path and os.path.exists(fon_path) and HAS_PIL:
                self.background_image = Image.open(fon_path)
                print("✅ Фоновое изображение загружено")
            else:
                print("⚠️ Фоновое изображение не загружено")
                self.background_image = None
        except Exception as e:
            print(f"❌ Ошибка загрузки фона: {e}")
            self.background_image = None
    
    def draw_background(self):
        self.canvas.delete("background")
        
        if self.background_image and HAS_PIL:
            width = self.root.winfo_width()
            height = self.root.winfo_height()
            
            if width < 10 or height < 10:
                width = 1200
                height = 700
            
            img_width = self.background_image.width
            img_height = self.background_image.height
            
            if img_width == 0 or img_height == 0:
                self.canvas.create_rectangle(0, 0, width, height, fill='#1a1a2e', tags="background")
                return
            
            for x in range(0, width + img_width, img_width):
                for y in range(0, height + img_height, img_height):
                    tile = self.background_image.copy()
                    photo = ImageTk.PhotoImage(tile)
                    self.background_tiles.append(photo)
                    self.canvas.create_image(x, y, image=photo, anchor='nw', tags="background")
        else:
            width = self.root.winfo_width()
            height = self.root.winfo_height()
            if width < 10 or height < 10:
                width = 1200
                height = 700
            self.canvas.create_rectangle(0, 0, width, height, fill='#1a1a2e', tags="background")
        
        self.canvas.tag_raise("widgets")
    
    def on_resize(self, event):
        if event.widget == self.root:
            self.draw_background()

    def on_layout_resize(self, event):
        if self._layout_resize_job:
            self.root.after_cancel(self._layout_resize_job)
        self._layout_resize_job = self.root.after(60, self._apply_layout, event.width, event.height)

    def _apply_layout(self, canvas_w, canvas_h):
        self._layout_resize_job = None

        m = self.LAYOUT_MARGIN
        content_w = max(canvas_w - m * 2, self.LAYOUT_MIN_LIST_W + self.LAYOUT_MIN_DESC_W + self.LAYOUT_GAP)

        list_w = int(content_w * self.LAYOUT_LIST_RATIO)
        list_w = max(list_w, self.LAYOUT_MIN_LIST_W)
        desc_w = content_w - list_w - self.LAYOUT_GAP
        if desc_w < self.LAYOUT_MIN_DESC_W:
            desc_w = self.LAYOUT_MIN_DESC_W
            list_w = content_w - desc_w - self.LAYOUT_GAP

        panel_h = max(canvas_h - 145 - m, self.LAYOUT_MIN_PANEL_H)

        self.canvas.itemconfig(self.header_window, width=content_w)
        self.canvas.itemconfig(self.btn_window, width=content_w)

        self.canvas.itemconfig(self.list_window, width=list_w, height=panel_h)

        desc_x = m + list_w + self.LAYOUT_GAP
        self.canvas.coords(self.status_window, desc_x, 115)
        self.canvas.coords(self.desc_window, desc_x, 145)
        self.canvas.itemconfig(self.desc_window, width=desc_w, height=panel_h)
    
    def start_title_scroll(self):
        text = self.game_title_label.cget('text')
        if len(text) > 25:
            text = text + '    ' + text + '    ' + text
            self.title_scroll_text = text
            self.title_scroll_pos = 0
            self.scroll_title()
    
    def scroll_title(self):
        if not hasattr(self, 'title_scroll_text') or not self.title_scroll_text:
            return
        
        text = self.title_scroll_text
        pos = self.title_scroll_pos
        display_text = text[pos:pos+30]
        
        if display_text:
            self.game_title_label.config(text=display_text)
        
        self.title_scroll_pos += 1
        if self.title_scroll_pos > len(text) - 30:
            self.title_scroll_pos = 0
        
        if self.title_scroll_timer:
            try:
                self.root.after_cancel(self.title_scroll_timer)
            except:
                pass
        
        self.title_scroll_timer = self.root.after(200, self.scroll_title)
    
    def stop_title_scroll(self):
        if self.title_scroll_timer:
            try:
                self.root.after_cancel(self.title_scroll_timer)
            except:
                pass
            self.title_scroll_timer = None
        self.title_scroll_text = None
    
    def create_widgets(self):
        self.LAYOUT_MARGIN = 20
        self.LAYOUT_LIST_RATIO = 0.35
        self.LAYOUT_GAP = 10
        self.LAYOUT_MIN_LIST_W = 260
        self.LAYOUT_MIN_DESC_W = 320
        self.LAYOUT_MIN_PANEL_H = 220

        init_content_w = 1200 - self.LAYOUT_MARGIN * 2

        # === ВЕРХНЯЯ ПАНЕЛЬ ===
        header = tk.Frame(self.canvas, bg='#1a1a2e')
        self.header_window = self.canvas.create_window(20, 15, anchor='nw',
                                                         width=init_content_w,
                                                         window=header, tags="widgets")
        
        title = tk.Label(header, text="🎮 Игры Switch", 
                         font=('Segoe UI', 20, 'bold'), 
                         fg='#ffffff', bg='#1a1a2e')
        title.pack(side='left')
        
        settings_btn = tk.Button(header, text="⚙️ Настройки", 
                                 font=('Segoe UI', 10),
                                 bg='#3a3a55', fg='#aaaaaa',
                                 padx=12, pady=4,
                                 relief='flat', cursor='hand2',
                                 command=self.show_settings)
        settings_btn.pack(side='right')
        
        # === КНОПКИ ===
        btn_frame = tk.Frame(self.canvas, bg='#1a1a2e')
        self.btn_window = self.canvas.create_window(20, 65, anchor='nw',
                                                      width=init_content_w,
                                                      window=btn_frame, tags="widgets")
        
        self.refresh_btn = tk.Button(btn_frame, text="🔄 Обновить", 
                                     font=('Segoe UI', 11),
                                     bg='#3a3a55', fg='white', 
                                     padx=20, pady=6,
                                     relief='flat', cursor='hand2',
                                     command=self.refresh)
        self.refresh_btn.pack(side='left')
        
        self.path_label = tk.Label(btn_frame, 
                                   text="📁 Папка не выбрана", 
                                   font=('Segoe UI', 10),
                                   fg='#888888', bg='#1a1a2e')
        self.path_label.pack(side='right', padx=10)
        
        # === ИНФОРМАЦИОННАЯ СТРОКА ===
        info = tk.Label(self.canvas, 
                        text="📂 Нажмите на игру → описание появится справа", 
                        font=('Segoe UI', 11),
                        fg='#aaaaaa', bg='#1a1a2e')
        info_window = self.canvas.create_window(20, 115, anchor='nw', window=info, tags="widgets")
        
        # === ЛЕВАЯ ПАНЕЛЬ: СПИСОК ИГР ===
        list_frame = tk.Frame(self.canvas, bg='#2d2d44')
        list_frame.pack_propagate(False)
        
        init_width = int(init_content_w * self.LAYOUT_LIST_RATIO)
        init_height = 500
        
        self.list_window = self.canvas.create_window(20, 145, anchor='nw', 
                                                 width=init_width, height=init_height,
                                                 window=list_frame, tags="widgets")
        
        # === СЧЁТЧИК И СТАТУС ===
        status_frame = tk.Frame(self.canvas, bg='#1a1a2e')
        self.status_window = self.canvas.create_window(init_width + 30, 115, anchor='nw', 
                                                   window=status_frame, tags="widgets")
        
        self.count_label = tk.Label(status_frame, 
                                    text="0 игр", 
                                    font=('Segoe UI', 11, 'bold'),
                                    fg='#8aadff', bg='#1a1a2e')
        self.count_label.pack(side='left', padx=(0, 15))
        
        self.status_label = tk.Label(status_frame, 
                                     text="Загрузка...", 
                                     font=('Segoe UI', 10),
                                     fg='#666666', bg='#1a1a2e')
        self.status_label.pack(side='left')
        
        # Заголовок списка с легендой
        legend_frame = tk.Frame(list_frame, bg='#2d2d44')
        legend_frame.pack(fill='x', pady=(5, 5))
        
        legend_label = tk.Label(legend_frame, text="📋 Игры", 
                                font=('Segoe UI', 12, 'bold'),
                                fg='#8aadff', bg='#2d2d44',
                                cursor='hand2')
        legend_label.pack(side='left')
        
        legend_indicators = tk.Frame(legend_frame, bg='#2d2d44')
        legend_indicators.pack(side='left', padx=(10, 0))
        
        for color in ['#4a8af4', '#4aca4a', '#e8a830']:
            sq_frame = tk.Frame(legend_indicators, bg='#2d2d44')
            sq_frame.pack(side='left', padx=2)
            sq = tk.Frame(sq_frame, bg=color, width=10, height=10, bd=1, relief='solid')
            sq.pack()
            sq.pack_propagate(False)
        
        def show_legend_tip(event):
            tip = tk.Toplevel(legend_frame)
            tip.overrideredirect(True)
            tip.configure(bg='#1a1a2e', bd=1, relief='solid')
            
            text = """🔵 = описание (клик для редактирования)
🟢 = обложка (клик — добавить фото/видео)
🟡 = геймплей (клик — добавить фото/видео)"""
            
            lbl = tk.Label(tip, text=text, 
                          font=('Segoe UI', 9),
                          fg='#cccccc', bg='#1a1a2e',
                          padx=10, pady=6, justify='left')
            lbl.pack()
            
            x = legend_frame.winfo_rootx()
            y = legend_frame.winfo_rooty() + 25
            tip.wm_geometry(f"+{x}+{y}")
            legend_frame.tip = tip
        
        def hide_legend_tip(event):
            if hasattr(legend_frame, 'tip'):
                legend_frame.tip.destroy()
                del legend_frame.tip
        
        legend_label.bind('<Enter>', show_legend_tip)
        legend_label.bind('<Leave>', hide_legend_tip)
        legend_indicators.bind('<Enter>', show_legend_tip)
        legend_indicators.bind('<Leave>', hide_legend_tip)
        
        # Canvas для списка
        self.list_canvas = tk.Canvas(list_frame, bg='#2d2d44', highlightthickness=0)
        scrollbar = tk.Scrollbar(list_frame, orient='vertical', command=self.list_canvas.yview)
        self.scrollable_frame = tk.Frame(self.list_canvas, bg='#2d2d44')
        
        self.scrollable_frame.bind(
            "<Configure>",
            lambda e: self.list_canvas.configure(scrollregion=self.list_canvas.bbox("all"))
        )
        
        self.list_canvas_window = self.list_canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.list_canvas.configure(yscrollcommand=scrollbar.set)

        def _on_list_canvas_configure(event):
            self.list_canvas.itemconfig(self.list_canvas_window, width=event.width)
        self.list_canvas.bind('<Configure>', _on_list_canvas_configure)
        
        self.list_canvas.pack(side='left', fill='both', expand=True)
        scrollbar.pack(side='right', fill='y')
        
        def on_list_mousewheel(event):
            self.list_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
            return "break"
        
        self.list_canvas.bind('<MouseWheel>', on_list_mousewheel)
        self.scrollable_frame.bind('<MouseWheel>', on_list_mousewheel)
        
        # === ПРАВАЯ ПАНЕЛЬ: ОПИСАНИЕ ===
        desc_frame = tk.Frame(self.canvas, bg='#2d2d44')
        desc_frame.pack_propagate(False)
        
        desc_width = init_content_w - init_width - self.LAYOUT_GAP
        desc_height = 500
        
        self.desc_window = self.canvas.create_window(init_width + 30, 145, anchor='nw',
                                                  width=desc_width, height=desc_height,
                                                  window=desc_frame, tags="widgets")
        self.desc_frame = desc_frame
        
        # === ВЕРХНЯЯ СТРОКА: НАЗВАНИЕ + КНОПКА ===
        top_row = tk.Frame(desc_frame, bg='#2d2d44')
        top_row.pack(fill='x', pady=(5, 5), padx=(10, 10))
        
        self.game_title_label = tk.Label(top_row, 
                                         text="Выберите игру", 
                                         font=('Segoe UI', 14, 'bold'),
                                         fg='#8aadff', bg='#2d2d44')
        self.game_title_label.pack(side='left')
        
        self.open_folder_btn = tk.Button(top_row, 
                                         text="📂 Открыть папку", 
                                         font=('Segoe UI', 10, 'bold'),
                                         bg='#4a6a8a', 
                                         fg='#ffffff',
                                         padx=12, 
                                         pady=4,
                                         relief='flat',
                                         cursor='hand2',
                                         bd=0,
                                         state='disabled',
                                         command=self.open_current_folder)
        self.open_folder_btn.pack(side='right')
        
        def on_open_enter(e):
            if self.open_folder_btn['state'] != 'disabled':
                e.widget.configure(bg='#5a7a9a')
        def on_open_leave(e):
            if self.open_folder_btn['state'] != 'disabled':
                e.widget.configure(bg='#4a6a8a')
        
        self.open_folder_btn.bind('<Enter>', on_open_enter)
        self.open_folder_btn.bind('<Leave>', on_open_leave)
        
        # === Обложка + геймплей ===
        images_frame = tk.Frame(desc_frame, bg='#2d2d44')
        images_frame.pack(fill='x', pady=10)
        
        # Обложка
        cover_frame = tk.Frame(images_frame, bg='#1a1a2e', bd=1, relief='solid')
        cover_frame.pack(side='left', padx=(0, 5), fill='both', expand=True)
        
        cover_label = tk.Label(cover_frame, text="📖 Обложка", 
                               font=('Segoe UI', 10),
                               fg='#888888', bg='#1a1a2e')
        cover_label.pack(pady=(5, 2))
        
        self.cover_label = tk.Label(cover_frame, 
                                    bg='#1a1a2e',
                                    text="Обложка отсутствует",
                                    font=('Segoe UI', 9),
                                    fg='#888888',
                                    cursor='hand2')
        self.cover_label.pack(padx=5, pady=(0, 5))
        self.cover_label.bind('<Enter>', lambda e: self.on_media_enter('cover', e))
        self.cover_label.bind('<Leave>', lambda e: self.on_media_leave('cover', e))
        self.cover_label.bind('<MouseWheel>', lambda e: self.on_media_scroll('cover', e))
        self.cover_label.bind('<Button-1>', lambda e: self.on_media_click('cover', e))
        
        # Геймплей
        gameplay_frame = tk.Frame(images_frame, bg='#1a1a2e', bd=1, relief='solid')
        gameplay_frame.pack(side='right', padx=(5, 0), fill='both', expand=True)
        
        gameplay_label = tk.Label(gameplay_frame, text="🎮 Геймплей", 
                                  font=('Segoe UI', 10),
                                  fg='#888888', bg='#1a1a2e')
        gameplay_label.pack(pady=(5, 2))
        
        self.gameplay_label = tk.Label(gameplay_frame, 
                                       bg='#1a1a2e',
                                       text="Геймплей отсутствует",
                                       font=('Segoe UI', 9),
                                       fg='#888888',
                                       cursor='hand2')
        self.gameplay_label.pack(padx=5, pady=(0, 5))
        
        self.gameplay_label.bind('<Enter>', lambda e: self.on_media_enter('gameplay', e))
        self.gameplay_label.bind('<Leave>', lambda e: self.on_media_leave('gameplay', e))
        self.gameplay_label.bind('<MouseWheel>', lambda e: self.on_media_scroll('gameplay', e))
        self.gameplay_label.bind('<Button-1>', lambda e: self.on_media_click('gameplay', e))
        
        # Текст описания
        desc_text_frame = tk.Frame(desc_frame, bg='#2d2d44')
        desc_text_frame.pack(fill='both', expand=True, pady=(5, 5), padx=10)
        
        self.desc_text = tk.Text(desc_text_frame, 
                                 font=('Segoe UI', 10),
                                 fg='#eeeeee', bg='#2d2d44',
                                 wrap='word',
                                 height=10,
                                 relief='flat', highlightthickness=1,
                                 highlightcolor='#4a4a6a', highlightbackground='#4a4a6a',
                                 cursor='hand2')
        self.desc_text.pack(side='left', fill='both', expand=True)
        
        desc_scroll = tk.Scrollbar(desc_text_frame, orient='vertical', command=self.desc_text.yview)
        desc_scroll.pack(side='right', fill='y')
        self.desc_text.configure(yscrollcommand=desc_scroll.set)
        
        # === НАДПИСЬ "ОПИСАНИЕ ОТСУТСТВУЕТ" ===
        self.desc_status_label = tk.Label(desc_frame, 
                                          text="",
                                          font=('Segoe UI', 9),
                                          fg='#888888', bg='#2d2d44',
                                          cursor='hand2')
        self.desc_status_label.pack(anchor='w', padx=10, pady=(0, 5))
        
        def on_refresh_enter(e):
            e.widget.configure(bg='#4a4a6a')
        def on_refresh_leave(e):
            e.widget.configure(bg='#3a3a55')
        
        self.refresh_btn.bind('<Enter>', on_refresh_enter)
        self.refresh_btn.bind('<Leave>', on_refresh_leave)
        
        def on_settings_enter(e):
            e.widget.configure(bg='#5a5a7a', fg='#ffffff')
        def on_settings_leave(e):
            e.widget.configure(bg='#3a3a55', fg='#aaaaaa')
        
        settings_btn.bind('<Enter>', on_settings_enter)
        settings_btn.bind('<Leave>', on_settings_leave)
    
    # ------------------------------------------------------------
    # ПОДСКАЗКИ И РЕДАКТИРОВАНИЕ ОПИСАНИЯ
    # ------------------------------------------------------------
    def show_desc_tip(self, event):
        if not self.current_desc_path:
            self.show_info_tip(event, 
                "📝 Описание отсутствует\n\n"
                "Как добавить:\n"
                "1. Нажмите 🔵 синий индикатор\n"
                "2. Введите текст\n"
                "3. Нажмите 'Сохранить'")
    
    def show_cover_tip(self, event):
        if not self.current_media['cover']:
            self.show_info_tip(event, 
                "📖 Обложка отсутствует\n\n"
                "Как добавить:\n"
                "1. Нажмите 🟢 зелёный индикатор\n"
                "2. Выберите фото или видео\n"
                "3. Можно добавить несколько — колесо мыши переключает\n"
                "4. Клик по видео открывает его в плеере\n"
                "   (для изображений — в просмотрщике)")
    
    def show_gameplay_tip(self, event):
        if not self.current_media['gameplay']:
            self.show_info_tip(event, 
                "🎮 Геймплей отсутствует\n\n"
                "Как добавить:\n"
                "1. Нажмите 🟡 жёлтый индикатор\n"
                "2. Выберите фото или видео\n"
                "3. Можно добавить несколько — колесо мыши переключает\n"
                "4. Клик по видео открывает его в плеере\n"
                "   (для изображений — в просмотрщике)")
    
    def show_info_tip(self, event, text):
        widget = event.widget
        tip = tk.Toplevel(widget)
        tip.overrideredirect(True)
        tip.configure(bg='#1a1a2e', bd=2, relief='solid')
        
        lbl = tk.Label(tip, text=text, 
                      font=('Segoe UI', 9),
                      fg='#cccccc', bg='#1a1a2e',
                      padx=12, pady=10, justify='left')
        lbl.pack()
        
        x = widget.winfo_rootx()
        y = widget.winfo_rooty() - 80
        tip.wm_geometry(f"+{x}+{y}")
        widget.tip = tip
    
    def hide_tip(self, event):
        if hasattr(event.widget, 'tip'):
            event.widget.tip.destroy()
            del event.widget.tip
    
    def edit_description(self, game):
        if not game:
            return
        
        game_path = os.path.join(self.switch_path, game)
        desc_path = os.path.join(game_path, 'Description.txt')
        
        if not os.path.exists(desc_path):
            try:
                with open(desc_path, 'w', encoding='utf-8') as f:
                    f.write('')
            except:
                messagebox.showerror("Ошибка", "Не удалось создать файл описания")
                return
        
        try:
            with open(desc_path, 'r', encoding='utf-8') as f:
                current_text = f.read()
        except:
            current_text = ''
        
        edit_window = tk.Toplevel(self.root)
        edit_window.title(f"Редактировать описание - {game}")
        edit_window.geometry("600x500")
        edit_window.configure(bg='#2d2d44')
        edit_window.resizable(True, True)
        
        try:
            icon_path = get_icon_path()
            if icon_path and os.path.exists(icon_path):
                edit_window.iconbitmap(icon_path)
        except:
            pass
        
        main_frame = tk.Frame(edit_window, bg='#2d2d44')
        main_frame.pack(fill='both', expand=True, padx=15, pady=15)
        
        tk.Label(main_frame, 
                 text=f"📝 Редактирование описания: {game}", 
                 font=('Segoe UI', 14, 'bold'),
                 fg='#ffffff', bg='#2d2d44').pack(pady=(0, 10), anchor='w')
        
        text_area = scrolledtext.ScrolledText(main_frame, 
                                              font=('Segoe UI', 10),
                                              fg='#eeeeee', bg='#1a1a2e',
                                              wrap='word',
                                              height=15,
                                              relief='flat')
        text_area.pack(fill='both', expand=True, pady=(0, 10))
        text_area.insert('1.0', current_text)
        
        btn_frame = tk.Frame(main_frame, bg='#2d2d44')
        btn_frame.pack(fill='x')
        
        def save_description():
            new_text = text_area.get('1.0', tk.END).strip()
            try:
                with open(desc_path, 'w', encoding='utf-8') as f:
                    f.write(new_text)
                messagebox.showinfo("Успех", "Описание сохранено!")
                edit_window.destroy()
                self.current_desc_path = desc_path
                self.load_games()
                self.show_game_description(game, desc_path, 
                                          self.current_media['cover'], self.current_media['gameplay'],
                                          delay=0)
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось сохранить: {e}")
        
        tk.Button(btn_frame, 
                  text="✅ Сохранить", 
                  font=('Segoe UI', 11),
                  bg='#4a8a4a', fg='white',
                  padx=20, pady=6,
                  relief='flat', cursor='hand2',
                  command=save_description).pack(side='left', padx=(0, 10))
        
        tk.Button(btn_frame, 
                  text="❌ Отмена", 
                  font=('Segoe UI', 11),
                  bg='#3a3a55', fg='#aaaaaa',
                  padx=20, pady=6,
                  relief='flat', cursor='hand2',
                  command=edit_window.destroy).pack(side='left')
    
    # ------------------------------------------------------------
    # РАБОТА С МЕДИА-ФАЙЛАМИ
    # ------------------------------------------------------------
    def _add_media_file(self, prefix, dialog_title):
        file_path = filedialog.askopenfilename(
            title=dialog_title,
            filetypes=[
                ("Фото и видео", "*.jpg *.jpeg *.png *.bmp *.gif *.webp *.mp4 *.mkv *.avi *.mov *.webm *.wmv *.m4v"),
                ("Изображения", "*.jpg *.jpeg *.png *.bmp *.gif *.webp"),
                ("Видео", "*.mp4 *.mkv *.avi *.mov *.webm *.wmv *.m4v"),
            ]
        )
        if not file_path:
            return None
        return file_path

    def choose_cover(self, game):
        self.add_cover(game)

    def add_cover(self, game):
        if not game:
            return
        
        file_path = self._add_media_file('cover', "Выберите обложку (фото или видео)")
        if not file_path:
            return
        
        game_path = os.path.join(self.switch_path, game)
        new_path = self._copy_as_next_numbered(file_path, game_path, 'cover')
        if not new_path:
            return
        
        messagebox.showinfo("Успех", f"Обложка добавлена: {os.path.basename(new_path)}")
        self.load_games()
        _, cover_media, gameplay_media = self.find_game_files(game_path)
        self.show_game_description(game, self.current_desc_path, cover_media, gameplay_media, delay=0)
    
    def add_gameplay(self, game):
        if not game:
            return
        
        file_path = self._add_media_file('game', "Выберите геймплей (фото или видео)")
        if not file_path:
            return
        
        game_path = os.path.join(self.switch_path, game)
        new_path = self._copy_as_next_numbered(file_path, game_path, 'game')
        if not new_path:
            return
        
        messagebox.showinfo("Успех", f"Геймплей добавлен: {os.path.basename(new_path)}")
        self.load_games()
        _, cover_media, gameplay_media = self.find_game_files(game_path)
        self.show_game_description(game, self.current_desc_path, cover_media, gameplay_media, delay=0)
    
    def _copy_as_next_numbered(self, file_path, game_path, prefix):
        existing = []
        try:
            for f in os.listdir(game_path):
                if f.lower().startswith(prefix) and self.is_media_file(f):
                    try:
                        num = int(''.join(filter(str.isdigit, f)))
                        existing.append(num)
                    except:
                        pass
        except Exception:
            pass
        
        next_num = max(existing) + 1 if existing else 1
        ext = os.path.splitext(file_path)[1]
        new_path = os.path.join(game_path, f'{prefix}{next_num}{ext}')
        
        try:
            shutil.copy2(file_path, new_path)
            return new_path
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось добавить файл: {e}")
            return None
    
    def is_image_file(self, filename):
        return filename.lower().endswith(IMAGE_EXTENSIONS)
    
    def is_video_file(self, filename):
        return filename.lower().endswith(VIDEO_EXTENSIONS)
    
    def is_media_file(self, filename):
        return self.is_image_file(filename) or self.is_video_file(filename)
    
    def find_game_files(self, game_path):
        desc_file = None
        cover_paths = []
        gameplay_paths = []
        
        try:
            for file in os.listdir(game_path):
                file_lower = file.lower()
                full_path = os.path.join(game_path, file)
                
                if file.startswith('.'):
                    continue
                
                if file_lower == 'description.txt':
                    desc_file = full_path
                elif file_lower.startswith('cover') and self.is_media_file(file):
                    cover_paths.append(full_path)
                elif file_lower.startswith('game') and self.is_media_file(file):
                    gameplay_paths.append(full_path)
        except:
            pass
        
        cover_paths.sort()
        gameplay_paths.sort()
        
        cover_media = [{'path': p, 'is_video': self.is_video_file(p)} for p in cover_paths]
        gameplay_media = [{'path': p, 'is_video': self.is_video_file(p)} for p in gameplay_paths]
        
        return desc_file, cover_media, gameplay_media
    
    def read_text_file(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except:
            try:
                with open(file_path, 'r', encoding='cp1251') as f:
                    return f.read()
            except:
                return "❌ Не удалось прочитать файл"
    
    def load_image(self, file_path, max_width=250, max_height=200):
        if not HAS_PIL or not file_path or not os.path.exists(file_path):
            return None
        
        try:
            img = Image.open(file_path)
            if img.mode in ('RGBA', 'LA', 'P'):
                try:
                    background = Image.new('RGB', img.size, (26, 26, 46))
                    if img.mode == 'P':
                        img = img.convert('RGBA')
                    background.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
                    img = background
                except:
                    img = img.convert('RGB')
            elif img.mode != 'RGB':
                img = img.convert('RGB')
            
            img.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)
            return ImageTk.PhotoImage(img)
        except Exception as e:
            print(f"Ошибка загрузки {file_path}: {e}")
            return None
    
    # ------------------------------------------------------------
    # НАСТРОЙКИ
    # ------------------------------------------------------------
    def show_settings(self):
        print("⚙️ Открываем окно настроек")
        
        settings_window = tk.Toplevel(self.root)
        settings_window.title("Настройки")
        settings_window.geometry("500x600")
        settings_window.configure(bg='#2d2d44')
        settings_window.resizable(False, False)
        
        settings_window.attributes('-topmost', True)
        settings_window.focus_force()
        settings_window.grab_set()
        
        try:
            icon_path = get_icon_path()
            if icon_path and os.path.exists(icon_path):
                settings_window.iconbitmap(icon_path)
        except:
            pass
        
        main_frame = tk.Frame(settings_window, bg='#2d2d44')
        main_frame.pack(fill='both', expand=True, padx=20, pady=20)
        
        tk.Label(main_frame, 
                 text="⚙️ Настройки", 
                 font=('Segoe UI', 18, 'bold'),
                 fg='#ffffff', bg='#2d2d44').pack(pady=(0, 15))
        
        tk.Label(main_frame, 
                 text="📁 Папка с играми:", 
                 font=('Segoe UI', 11),
                 fg='#cccccc', bg='#2d2d44').pack(anchor='w')
        
        path_var = tk.StringVar(value=self.switch_path if self.switch_path else "Не выбрана")
        path_entry = tk.Entry(main_frame, 
                              textvariable=path_var,
                              font=('Segoe UI', 10),
                              bg='#ffffff',
                              fg='#000000',
                              relief='solid',
                              bd=1)
        path_entry.pack(fill='x', pady=(2, 5))
        
        def choose_folder():
            settings_window.attributes('-topmost', False)
            settings_window.grab_release()
            
            folder = filedialog.askdirectory(title="Выберите папку Switch с играми")
            
            settings_window.attributes('-topmost', True)
            settings_window.grab_set()
            settings_window.focus_force()
            
            if folder:
                self.switch_path = folder
                path_var.set(folder)
                self.path_label.config(text=f"📁 {folder}")
                self.config['Settings']['last_path'] = folder
                self.save_config()
                print(f"📁 Выбрана папка: {folder}")
            else:
                print("📁 Выбор папки отменён")
        
        tk.Button(main_frame, 
                  text="📂 Выбрать папку", 
                  font=('Segoe UI', 10),
                  bg='#4a4a6a', fg='white',
                  padx=15, pady=4,
                  relief='flat', cursor='hand2',
                  command=choose_folder).pack(anchor='w')
        
        tk.Frame(main_frame, bg='#4a4a6a', height=1).pack(fill='x', pady=10)
        
        auto_scan_var = tk.BooleanVar(value=self.config.getboolean('Settings', 'auto_scan', fallback=True))
        
        def on_auto_changed():
            self.config['Settings']['auto_scan'] = str(auto_scan_var.get())
            self.save_config()
        
        tk.Checkbutton(main_frame, 
                       text="🔍 Авто-сканирование при запуске", 
                       variable=auto_scan_var,
                       font=('Segoe UI', 11),
                       fg='#cccccc', bg='#2d2d44',
                       selectcolor='#4a4a6a',
                       activebackground='#2d2d44',
                       activeforeground='#ffffff',
                       command=on_auto_changed).pack(anchor='w')
        
        tk.Frame(main_frame, bg='#4a4a6a', height=1).pack(fill='x', pady=10)
        
        tk.Label(main_frame, 
                 text="🔄 Частота смены обложек и геймплей-картинок (мс):", 
                 font=('Segoe UI', 11),
                 fg='#cccccc', bg='#2d2d44').pack(anchor='w')
        
        interval_var = tk.StringVar(value=str(self.slide_interval))
        interval_entry = tk.Entry(main_frame, 
                                  textvariable=interval_var,
                                  font=('Segoe UI', 11),
                                  bg='#ffffff',
                                  fg='#000000',
                                  width=10,
                                  relief='solid',
                                  bd=1)
        interval_entry.pack(anchor='w', pady=(2, 0))
        
        tk.Label(main_frame, 
                 text="(500 - 10000 мс)", 
                 font=('Segoe UI', 9),
                 fg='#666666', bg='#2d2d44').pack(anchor='w')
        
        play_video_full_var = tk.BooleanVar(value=self.play_video_full)
        
        def on_video_full_changed():
            self.config['Settings']['play_video_full'] = str(play_video_full_var.get())
            self.save_config()
        
        video_cb = tk.Checkbutton(main_frame, 
                                  text="▶️ Воспроизводить видео до конца (игнорировать таймер слайд-шоу)", 
                                  variable=play_video_full_var,
                                  font=('Segoe UI', 11),
                                  fg='#cccccc', bg='#2d2d44',
                                  selectcolor='#4a4a6a',
                                  activebackground='#2d2d44',
                                  activeforeground='#ffffff',
                                  wraplength=380,
                                  justify='left',
                                  command=on_video_full_changed)
        video_cb.pack(anchor='w', pady=(5, 0))
        
        tk.Frame(main_frame, bg='#4a4a6a', height=1).pack(fill='x', pady=10)
        
        def show_about():
            print("ℹ️ Открываем окно Об авторе")
            about_window = tk.Toplevel(settings_window)
            about_window.title("Об авторе")
            about_window.geometry("450x550")
            about_window.configure(bg='#2d2d44')
            about_window.resizable(False, False)
            
            about_window.attributes('-topmost', True)
            about_window.focus_force()
            
            try:
                icon_path = get_icon_path()
                if icon_path and os.path.exists(icon_path):
                    about_window.iconbitmap(icon_path)
            except:
                pass
            
            main_frame2 = tk.Frame(about_window, bg='#2d2d44')
            main_frame2.pack(fill='both', expand=True, padx=25, pady=25)
            
            title_label = tk.Label(main_frame2, 
                                   text="ℹ️ Об авторе", 
                                   font=('Segoe UI', 20, 'bold'),
                                   fg='#ffffff', bg='#2d2d44')
            title_label.pack(pady=(0, 15))
            
            photo_frame = tk.Frame(main_frame2, bg='#1a1a2e', bd=2, relief='solid')
            photo_frame.pack(pady=10)
            
            photo_image = None
            photo_path = get_photo_path()
            if photo_path and os.path.exists(photo_path) and HAS_PIL:
                try:
                    img = Image.open(photo_path)
                    img.thumbnail((200, 250), Image.Resampling.LANCZOS)
                    photo_image = ImageTk.PhotoImage(img)
                except Exception as e:
                    print(f"Ошибка загрузки фото: {e}")
            
            if photo_image:
                photo_label = tk.Label(photo_frame, image=photo_image, bg='#1a1a2e')
                photo_label.image = photo_image
                photo_label.pack(padx=8, pady=8)
            else:
                placeholder = tk.Label(photo_frame, 
                                       text="📷\nФото\nне загружено", 
                                       font=('Segoe UI', 16),
                                       fg='#666666', bg='#1a1a2e',
                                       width=22, height=12)
                placeholder.pack(padx=8, pady=8)
            
            info_text = """Создал: Шувалов Павел Витальевич
Дата: 14.07.2026г. v_1.1

Программа для удобного просмотра
игр Nintendo Switch с описаниями."""
            
            info_label = tk.Label(main_frame2, 
                                  text=info_text,
                                  font=('Segoe UI', 12),
                                  fg='#cccccc', bg='#2d2d44',
                                  justify='center')
            info_label.pack(pady=15)
            
            close_btn = tk.Button(main_frame2, 
                                 text="✕ Закрыть", 
                                 font=('Segoe UI', 12, 'bold'),
                                 bg='#4a4a6a', 
                                 fg='#ffffff',
                                 width=12,
                                 height=1,
                                 padx=10,
                                 pady=8,
                                 relief='raised',
                                 cursor='hand2',
                                 bd=2,
                                 activebackground='#5a5a8a',
                                 activeforeground='#ffffff',
                                 command=about_window.destroy)
            close_btn.pack(pady=(10, 0))
            
            def on_btn_enter(e):
                e.widget.configure(bg='#5a5a8a', bd=3)
            def on_btn_leave(e):
                e.widget.configure(bg='#4a4a6a', bd=2)
            
            close_btn.bind('<Enter>', on_btn_enter)
            close_btn.bind('<Leave>', on_btn_leave)
            
            about_window.update_idletasks()
            width = about_window.winfo_width()
            height = about_window.winfo_height()
            x = (about_window.winfo_screenwidth() // 2) - (width // 2)
            y = (about_window.winfo_screenheight() // 2) - (height // 2)
            about_window.geometry(f"{width}x{height}+{x}+{y}")
            
            about_window.protocol("WM_DELETE_WINDOW", lambda: [about_window.destroy(), settings_window.focus_force()])
        
        tk.Button(main_frame, 
                  text="ℹ️ Об авторе", 
                  font=('Segoe UI', 11),
                  bg='#4a4a6a', fg='white',
                  padx=20, pady=6,
                  relief='flat', cursor='hand2',
                  command=show_about).pack(pady=5)
        
        tk.Frame(main_frame, bg='#4a4a6a', height=1).pack(fill='x', pady=10)
        
        btn_frame = tk.Frame(main_frame, bg='#2d2d44')
        btn_frame.pack(pady=10)
        
        def save_settings():
            try:
                new_interval = int(interval_var.get())
                if new_interval < 500 or new_interval > 10000:
                    messagebox.showwarning("Предупреждение", "Интервал должен быть от 500 до 10000 мс")
                    return
                
                self.slide_interval = new_interval
                self.config['Settings']['slide_interval'] = str(new_interval)
                self.play_video_full = play_video_full_var.get()
                self.config['Settings']['play_video_full'] = str(self.play_video_full)
                self.save_config()
                
                for role in ('cover', 'gameplay'):
                    if len(self.current_media.get(role, [])) > 1:
                        self.stop_slideshow(role)
                        self.start_slideshow(role)
                
                settings_window.destroy()
                self.load_games()
                messagebox.showinfo("Успех", "Настройки сохранены!")
            except ValueError:
                messagebox.showerror("Ошибка", "Введите корректное число")
        
        tk.Button(btn_frame, 
                  text="✅ Сохранить", 
                  font=('Segoe UI', 11),
                  bg='#4a8a4a', fg='white',
                  padx=20, pady=6,
                  relief='flat', cursor='hand2',
                  command=save_settings).pack(side='left', padx=5)
        
        def cancel_and_close():
            settings_window.destroy()
            self.load_games()
        
        tk.Button(btn_frame, 
                  text="❌ Отмена", 
                  font=('Segoe UI', 11),
                  bg='#3a3a55', fg='#aaaaaa',
                  padx=20, pady=6,
                  relief='flat', cursor='hand2',
                  command=cancel_and_close).pack(side='left', padx=5)
        
        settings_window.protocol("WM_DELETE_WINDOW", lambda: [settings_window.destroy(), self.load_games()])
        
        settings_window.update_idletasks()
        width = settings_window.winfo_width()
        height = settings_window.winfo_height()
        x = (settings_window.winfo_screenwidth() // 2) - (width // 2)
        y = (settings_window.winfo_screenheight() // 2) - (height // 2)
        settings_window.geometry(f"{width}x{height}+{x}+{y}")
    
    # ------------------------------------------------------------
    # МЕДИА СОБЫТИЯ (наведение, скролл, клик)
    # ------------------------------------------------------------
    def on_media_enter(self, role, event):
        self.slides[role]['is_hovering'] = True
        self.stop_slideshow(role)
        if role == 'cover':
            self.show_cover_tip(event)
        else:
            self.show_gameplay_tip(event)
    
    def on_media_leave(self, role, event):
        self.slides[role]['is_hovering'] = False
        if len(self.current_media.get(role, [])) > 1:
            self.start_slideshow(role)
        self.hide_tip(event)
    
    def on_media_scroll(self, role, event):
        items = self.current_media.get(role, [])
        if len(items) <= 1:
            return
        idx = self.slides[role]['index']
        if event.delta > 0:
            idx = (idx - 1) % len(items)
        else:
            idx = (idx + 1) % len(items)
        self.slides[role]['index'] = idx
        self.render_media(role, items[idx])
    
    def on_media_click(self, role, event):
        path = self._displayed_media_path.get(role)
        if path:
            try:
                os.startfile(path)
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{path}\n\n{str(e)}")
    
    # ------------------------------------------------------------
    # СЛАЙД-ШОУ
    # ------------------------------------------------------------
    def start_slideshow(self, role):
        items = self.current_media.get(role, [])
        if len(items) <= 1 or self.slides[role]['is_hovering']:
            return
        self.slides[role]['is_sliding'] = True
        self.next_slide(role)
    
    def next_slide(self, role):
        st = self.slides[role]
        items = self.current_media.get(role, [])
        if not st['is_sliding'] or st['is_hovering'] or len(items) <= 1:
            return
        
        if self.play_video_full:
            player = self.video_players.get(role)
            if player and player['running']:
                st['timer'] = self.root.after(self.slide_interval, self.next_slide, role)
                return
        
        st['index'] = (st['index'] + 1) % len(items)
        self.render_media(role, items[st['index']])
        st['timer'] = self.root.after(self.slide_interval, self.next_slide, role)
    
    def stop_slideshow(self, role):
        st = self.slides[role]
        st['is_sliding'] = False
        if st['timer']:
            self.root.after_cancel(st['timer'])
            st['timer'] = None
    
    # ------------------------------------------------------------
    # ОСНОВНАЯ ЛОГИКА ЗАГРУЗКИ И ОТОБРАЖЕНИЯ
    # ------------------------------------------------------------
    def auto_start(self):
        print("🚀 Автозапуск...")
        auto_scan = self.config.getboolean('Settings', 'auto_scan', fallback=True)
        last_path = self.config.get('Settings', 'last_path', fallback='')
        self.slide_interval = self.config.getint('Settings', 'slide_interval', fallback=3000)
        self.play_video_full = self.config.getboolean('Settings', 'play_video_full', fallback=False)
        
        if last_path and os.path.exists(last_path):
            self.switch_path = last_path
            self.path_label.config(text=f"📁 {last_path}")
            self.load_games()
            return
        
        if auto_scan:
            current_dir = os.path.dirname(os.path.abspath(__file__))
            switch_path = os.path.join(current_dir, 'Switch')
            if os.path.exists(switch_path) and os.path.isdir(switch_path):
                self.switch_path = switch_path
                self.path_label.config(text=f"📁 {switch_path}")
                self.config['Settings']['last_path'] = switch_path
                self.save_config()
                self.load_games()
                return
        
        self.path_label.config(text="📁 Папка не выбрана (настройки)")
        self.show_empty("😕 Папка не выбрана\n\nНажмите ⚙️ Настройки → Выбрать папку")
    
    def load_games(self):
        if self._update_job:
            self.root.after_cancel(self._update_job)
            self._update_job = None
        
        if not self.switch_path:
            self.show_empty("😕 Папка не выбрана\n\nНажмите ⚙️ Настройки → Выбрать папку")
            self.count_label.config(text="0 игр")
            self.status_label.config(text="⚠️ Папка не выбрана")
            return
        
        if not os.path.exists(self.switch_path):
            self.show_empty(f"😕 Папка не существует:\n{self.switch_path}")
            self.count_label.config(text="0 игр")
            self.status_label.config(text="⚠️ Папка не найдена")
            return
        
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        
        self.games = []
        self.status_label.config(text="⏳ Сканирование...")
        self.root.update()
        
        try:
            items = os.listdir(self.switch_path)
            folders = []
            for item in items:
                item_path = os.path.join(self.switch_path, item)
                if os.path.isdir(item_path) and not item.startswith('.'):
                    folders.append(item)
            
            if not folders:
                self.show_empty("😕 Папок с играми не найдено")
                self.count_label.config(text="0 игр")
                self.status_label.config(text="Готов")
                return
            
            self.games = sorted(folders)
            self.show_games()
            self.status_label.config(text=f"✅ Найдено {len(self.games)} игр")
            
            if self.selected_index >= len(self.games):
                self.selected_index = -1
            if self.selected_index == -1 and self.games:
                self.selected_index = 0
                self.select_game(0)
            elif self.selected_index >= 0:
                self.select_game(self.selected_index)
                
        except Exception as e:
            self.show_empty(f"❌ Ошибка: {str(e)}")
            self.count_label.config(text="0 игр")
            self.status_label.config(text="❌ Ошибка")
    
    def show_games(self):
        self.count_label.config(text=f"{len(self.games)} игр")
        
        for i, game in enumerate(self.games):
            frame = tk.Frame(self.scrollable_frame, bg='#3a3a55', bd=0, relief='flat')
            frame.pack(fill='x', pady=1, padx=2)
            frame.game_name = game
            frame.index = i

            def on_item_mousewheel(event, canvas=self.list_canvas):
                canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
                return "break"
            frame.bind('<MouseWheel>', on_item_mousewheel)

            indicators_left = tk.Frame(frame, bg='#3a3a55')
            indicators_left.pack(side='right', padx=(6, 4), pady=1)

            game_path = os.path.join(self.switch_path, game)
            desc_file, cover_media, gameplay_media = self.find_game_files(game_path)
            has_desc = desc_file is not None
            has_cover = len(cover_media) > 0
            has_gameplay = len(gameplay_media) > 0

            def create_indicator(parent, color, active, tooltip_text, action, game_name):
                ind_frame = tk.Frame(parent, bg=color if active else '#3d3d5c',
                                     width=12, height=12, bd=1, relief='solid')
                ind_frame.pack(side='left', padx=2)
                ind_frame.pack_propagate(False)

                dot = tk.Label(ind_frame, text="",
                               bg=color if active else '#3d3d5c')
                dot.pack(fill='both', expand=True)

                ind_frame.configure(cursor='hand2')

                def on_enter(e):
                    tip = tk.Toplevel(ind_frame)
                    tip.overrideredirect(True)
                    tip.configure(bg='#1a1a2e', bd=1, relief='solid')
                    status_text = "✅" if active else "❌"
                    lbl = tk.Label(tip, text=f"{tooltip_text} {status_text}\n(клик для редактирования)",
                                   font=('Segoe UI', 8),
                                   fg='#cccccc', bg='#1a1a2e',
                                   padx=8, pady=4, justify='left')
                    lbl.pack()
                    x = ind_frame.winfo_rootx()
                    y = ind_frame.winfo_rooty() - 38
                    tip.wm_geometry(f"+{x}+{y}")
                    ind_frame.tip = tip

                def on_leave(e):
                    if hasattr(ind_frame, 'tip'):
                        ind_frame.tip.destroy()
                        del ind_frame.tip

                def on_click(e):
                    action(game_name)

                ind_frame.bind('<Enter>', on_enter)
                ind_frame.bind('<Leave>', on_leave)
                ind_frame.bind('<Button-1>', on_click)
                dot.bind('<Button-1>', on_click)
                return ind_frame

            create_indicator(indicators_left, '#4a8af4', has_desc,
                            "📝 Описание", self.edit_description, game)
            create_indicator(indicators_left, '#4aca4a', has_cover,
                            "📖 Обложка", self.add_cover, game)
            create_indicator(indicators_left, '#e8a830', has_gameplay,
                            "🎮 Геймплей", self.add_gameplay, game)

            name_label = tk.Label(frame, text=self._truncate_name(game),
                                  font=('Segoe UI', 11),
                                  fg='#ffffff', bg='#3a3a55',
                                  cursor='hand2', anchor='w')
            name_label.pack(side='left', padx=(5, 0), pady=2, fill='x', expand=True)
            name_label.bind('<MouseWheel>', on_item_mousewheel)

            frame.desc_file = desc_file
            frame.cover_media = cover_media
            frame.gameplay_media = gameplay_media
            
            def on_click(e, g=game, idx=i, df=desc_file, cm=cover_media, gm=gameplay_media):
                self.selected_index = idx
                self.highlight_game(g)
                self.scroll_to_index(idx)
                self.show_game_description(g, df, cm, gm, delay=300)
            
            frame.bind('<Button-1>', on_click)
            name_label.bind('<Button-1>', on_click)
    
    def highlight_game(self, game_name):
        for child in self.scrollable_frame.winfo_children():
            if hasattr(child, 'game_name'):
                if child.game_name == game_name:
                    child.config(bg='#4a6a8a')
                else:
                    child.config(bg='#3a3a55')
    
    def show_game_description(self, game, desc_file, cover_media, gameplay_media, delay=300):
        if self._update_job:
            self.root.after_cancel(self._update_job)
            self._update_job = None
        
        if delay > 0:
            self._update_job = self.root.after(delay, self._do_show_game_description,
                                               game, desc_file, cover_media, gameplay_media)
        else:
            self._do_show_game_description(game, desc_file, cover_media, gameplay_media)
    
    def _do_show_game_description(self, game, desc_file, cover_media, gameplay_media):
        self._update_job = None
        
        self.current_game = game
        self.current_desc_path = desc_file
        self.current_media['cover'] = cover_media or []
        self.current_media['gameplay'] = gameplay_media or []
        
        self.stop_title_scroll()
        self._stop_all_video()
        
        self.game_title_label.config(text=game)
        self.open_folder_btn.config(state='normal')
        
        if len(game) > 25:
            self.root.after(500, self.start_title_scroll)
        else:
            self.game_title_label.config(text=game)
        
        for role in ('cover', 'gameplay'):
            self.stop_slideshow(role)
            self.slides[role]['index'] = 0
        
        cover_items = self.current_media['cover']
        self.render_media('cover', cover_items[0] if cover_items else None)
        if len(cover_items) > 1:
            self.start_slideshow('cover')
        
        gameplay_items = self.current_media['gameplay']
        self.render_media('gameplay', gameplay_items[0] if gameplay_items else None)
        if len(gameplay_items) > 1:
            self.start_slideshow('gameplay')
        
        if desc_file:
            text = self.read_text_file(desc_file)
            self.desc_text.delete('1.0', tk.END)
            self.desc_text.insert('1.0', text)
            self.desc_text.configure(fg='#eeeeee')
            self.desc_status_label.config(text="", cursor='arrow')
        else:
            self.desc_text.delete('1.0', tk.END)
            self.desc_text.insert('1.0', "📝 Описание отсутствует")
            self.desc_text.configure(fg='#888888')
            self.desc_status_label.config(text="📝 Описание отсутствует (нажмите 🔵 для добавления)", 
                                         fg='#888888', cursor='hand2')
            self.desc_status_label.bind('<Enter>', self.show_desc_tip)
            self.desc_status_label.bind('<Leave>', self.hide_tip)
        
        if game:
            self.highlight_game(game)
    
    def open_current_folder(self):
        if self.current_game:
            folder_path = os.path.join(self.switch_path, self.current_game)
            try:
                os.startfile(folder_path)
            except:
                messagebox.showinfo("Путь", f"Папка: {folder_path}")
    
    def show_empty(self, message):
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        label = tk.Label(self.scrollable_frame, text=message,
                         font=('Segoe UI', 14),
                         fg='#888888', bg='#2d2d44')
        label.pack(expand=True, pady=50)
    
    def refresh(self):
        if self.switch_path:
            self.load_games()
        else:
            self.show_empty("😕 Папка не выбрана\n\nНажмите ⚙️ Настройки → Выбрать папку")
            self.count_label.config(text="0 игр")
            self.status_label.config(text="⚠️ Папка не выбрана")


def main():
    print("="*50)
    print("🎮 Запуск Switch Menu")
    print("="*50)
    
    root = tk.Tk()
    app = SwitchMenu(root)
    root.mainloop()


if __name__ == "__main__":
    main()